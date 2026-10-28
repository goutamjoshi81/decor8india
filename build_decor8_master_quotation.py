import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
import win32com.client

def build_perfect_white_gold_quotation():
    output_excel_paths = [
        "c:/Users/gouta/Desktop/decor8india/public/DECOR8_INDIA_QUOTATION_MR_GOVINDA_D1304.xlsx",
        "c:/Users/gouta/Desktop/decor8india/DECOR8_INDIA_QUOTATION_MR_GOVINDA_D1304.xlsx",
        "c:/Users/gouta/Desktop/decor8india/DECOR8_INDIA_QUOTATION_WHITE_AND_GOLD.xlsx"
    ]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # ==============================================================================
    # --- LUXURY COLOR PALETTE (DECOR8 INDIA THEME: WHITE & GOLD) ---
    # ==============================================================================
    COLOR_GOLD_DARK     = "9E7D3B"   # Deep Warm Gold (Titles, Accent borders)
    COLOR_GOLD_METALLIC = "C5A059"   # Primary Metallic Gold (Badges, Key Totals)
    COLOR_GOLD_DEEP     = "8C7355"   # Warm Taupe Gold (Table Headers - match ref images)
    COLOR_GOLD_LIGHT    = "F5EEDB"   # Soft Champagne Cream (Subtotals, summary bars)
    COLOR_ZEBRA_LIGHT   = "FAF8F3"   # Ivory tint for alternate rows
    COLOR_CARD_BG       = "FDFBF7"   # Warm card background
    COLOR_BORDER_GOLD   = "D4AF37"   # Gold border
    COLOR_BORDER_SUBTLE = "EADCB5"   # Light subtle gold border
    COLOR_BORDER_ROW    = "EFE7D8"   # Table row border
    COLOR_WHITE         = "FFFFFF"   # Pure White
    COLOR_DARK          = "1A1A1A"   # Jet Charcoal (High contrast readability)
    COLOR_MUTED         = "555555"   # Subtitles, remarks
    COLOR_SUCCESS       = "1B6A3B"   # Green accent for "INCLUDED"
    COLOR_HEADER_BG     = "111317"   # Deep Obsidian for Master Title bar

    # Borders
    thin_gold   = Side(border_style="thin", color=COLOR_BORDER_GOLD)
    medium_gold = Side(border_style="medium", color=COLOR_BORDER_GOLD)
    double_gold = Side(border_style="double", color=COLOR_BORDER_GOLD)
    thin_row    = Side(border_style="thin", color=COLOR_BORDER_ROW)

    cell_border        = Border(left=thin_row, right=thin_row, top=thin_row, bottom=thin_row)
    header_border      = Border(left=thin_gold, right=thin_gold, top=medium_gold, bottom=medium_gold)
    subtotal_border    = Border(left=thin_gold, right=thin_gold, top=thin_gold, bottom=medium_gold)
    grand_total_border = Border(left=medium_gold, right=medium_gold, top=medium_gold, bottom=double_gold)

    # Fills
    fill_white       = PatternFill(start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid")
    fill_tbl_header  = PatternFill(start_color=COLOR_GOLD_DEEP, end_color=COLOR_GOLD_DEEP, fill_type="solid")
    fill_subtotal    = PatternFill(start_color=COLOR_GOLD_LIGHT, end_color=COLOR_GOLD_LIGHT, fill_type="solid")
    fill_zebra       = PatternFill(start_color=COLOR_ZEBRA_LIGHT, end_color=COLOR_ZEBRA_LIGHT, fill_type="solid")
    fill_grand_total = PatternFill(start_color=COLOR_GOLD_DEEP, end_color=COLOR_GOLD_DEEP, fill_type="solid")
    fill_card        = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type="solid")
    fill_gold_bar    = PatternFill(start_color=COLOR_GOLD_METALLIC, end_color=COLOR_GOLD_METALLIC, fill_type="solid")
    fill_meta_lbl    = PatternFill(start_color="EFE6D0", end_color="EFE6D0", fill_type="solid")

    # Fonts
    font_main_title  = Font(name="Segoe UI", size=14, bold=True, color=COLOR_WHITE)
    font_brand_title = Font(name="Segoe UI", size=15, bold=True, color=COLOR_DARK)
    font_brand_sub   = Font(name="Segoe UI", size=8.5, bold=False, color=COLOR_MUTED)
    font_brand_rt    = Font(name="Georgia", size=15, bold=True, color=COLOR_GOLD_DARK)
    font_tagline_rt  = Font(name="Segoe UI", size=8.5, italic=True, color=COLOR_MUTED)

    font_client_hdr  = Font(name="Segoe UI", size=10, bold=True, color=COLOR_DARK)
    font_sec_badge   = Font(name="Segoe UI", size=8.5, bold=True, color=COLOR_GOLD_DARK)
    font_sec_title   = Font(name="Segoe UI", size=13, bold=True, color=COLOR_DARK)
    font_sec_sub     = Font(name="Segoe UI", size=9, italic=True, color=COLOR_MUTED)
    font_sec_total   = Font(name="Segoe UI", size=13, bold=True, color=COLOR_DARK)

    font_tbl_header  = Font(name="Segoe UI", size=9, bold=True, color=COLOR_WHITE)
    font_cell        = Font(name="Segoe UI", size=9, bold=False, color=COLOR_DARK)
    font_cell_bold   = Font(name="Segoe UI", size=9, bold=True, color=COLOR_DARK)
    font_cell_amount = Font(name="Segoe UI", size=9.5, bold=True, color=COLOR_DARK)
    font_cell_subtot = Font(name="Segoe UI", size=10, bold=True, color=COLOR_DARK)
    font_bullet      = Font(name="Segoe UI", size=8.5, bold=False, color="333333")
    font_bullet_hdr  = Font(name="Segoe UI", size=9, bold=True, color=COLOR_DARK)
    font_included    = Font(name="Segoe UI", size=8.5, bold=True, color=COLOR_SUCCESS)

    font_grand_lbl   = Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE)
    font_grand_val   = Font(name="Segoe UI", size=13, bold=True, color=COLOR_WHITE)

    font_terms_hdr   = Font(name="Segoe UI", size=9.5, bold=True, color=COLOR_DARK)
    font_terms_item  = Font(name="Segoe UI", size=8, bold=False, color="222222")

    # Alignments
    align_left       = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    align_left_top   = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True)
    align_right      = Alignment(horizontal="right", vertical="center", indent=1)
    align_center     = Alignment(horizontal="center", vertical="center")
    align_center_top = Alignment(horizontal="center", vertical="top")

    FMT_CURRENCY = '[$₹-en-IN] #,##,##0'

    def apply_style(ws, start_col, start_row, end_col, end_row, font=None, fill=None, border=None, alignment=None):
        for r_idx in range(start_row, end_row + 1):
            for c_idx in range(start_col, end_col + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                if font: cell.font = font
                if fill: cell.fill = fill
                if border: cell.border = border
                if alignment: cell.alignment = alignment

    # ==============================================================================
    # --- SHEET 1: MASTER LUXURY PROPOSAL (PERFECT 6-COLUMN GRID A TO F) ---
    # ==============================================================================
    ws1 = wb.create_sheet(title="Luxury Proposal (17 Sections)")
    ws1.views.sheetView[0].showGridLines = True

    # Column Widths across EXACTLY 6 columns (A to F):
    # Total characters = 20 + 36 + 7 + 7 + 12 + 15 = 97. Fits A4 portrait perfectly!
    col_widths_1 = {
        'A': 20,    # Item Name / Category (plenty of room for labels)
        'B': 36,    # Specification & Materials
        'C': 7,     # Qty
        'D': 7,     # Unit
        'E': 12,    # Rate (₹)
        'F': 15     # Amount (₹)
    }
    for col, width in col_widths_1.items():
        ws1.column_dimensions[col].width = width

    # Page Setup in openpyxl
    ws1.page_setup.orientation = ws1.ORIENTATION_PORTRAIT
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    ws1.sheet_properties.pageSetUpPr.fitToPage = True
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 0
    ws1.page_margins.left = 0.3
    ws1.page_margins.right = 0.3
    ws1.page_margins.top = 0.25
    ws1.page_margins.bottom = 0.25

    r = 1

    # 1. Top Luxury Gold Bar
    ws1.row_dimensions[r].height = 4
    apply_style(ws1, 1, r, 6, r, fill=fill_gold_bar)
    r += 1

    # 2. Main Title Banner
    ws1.row_dimensions[r].height = 24
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "LUXURY INTERIOR DESIGN & TURNKEY FITOUT QUOTATION"
    apply_style(ws1, 1, r, 6, r, font=font_main_title, fill=PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid"), alignment=align_center)
    r += 1

    # 3. Company Info & Brand Block
    ws1.row_dimensions[r].height = 18
    ws1.row_dimensions[r+1].height = 13
    ws1.row_dimensions[r+2].height = 13
    ws1.row_dimensions[r+3].height = 13
    ws1.row_dimensions[r+4].height = 14

    # Left: Company Info (Columns A to C)
    ws1.merge_cells(f"A{r}:C{r}")
    ws1[f"A{r}"] = "DECOR8 INDIA"
    ws1[f"A{r}"].font = font_brand_title
    ws1[f"A{r}"].alignment = align_left

    ws1.merge_cells(f"A{r+1}:C{r+1}")
    ws1[f"A{r+1}"] = "#14, Sy NO 36/1 Vasanth Vallabhnagar, Vasanthapura"
    ws1[f"A{r+1}"].font = font_brand_sub
    ws1[f"A{r+1}"].alignment = align_left

    ws1.merge_cells(f"A{r+2}:C{r+2}")
    ws1[f"A{r+2}"] = "Uttarahalli Hobli, Bengaluru - 560061"
    ws1[f"A{r+2}"].font = font_brand_sub
    ws1[f"A{r+2}"].alignment = align_left

    ws1.merge_cells(f"A{r+3}:C{r+3}")
    ws1[f"A{r+3}"] = "+91 88841 31414 , +91 93805 23743"
    ws1[f"A{r+3}"].font = font_brand_sub
    ws1[f"A{r+3}"].alignment = align_left

    ws1.merge_cells(f"A{r+4}:C{r+4}")
    ws1[f"A{r+4}"] = "GSTIN: 29CHPPB0944C2ZD"
    ws1[f"A{r+4}"].font = font_cell_bold
    ws1[f"A{r+4}"].alignment = align_left

    # Right: Brand Identity Block (Columns D to F)
    ws1.merge_cells(f"D{r}:F{r+1}")
    ws1[f"D{r}"] = "DECOR8 INDIA"
    ws1[f"D{r}"].font = font_brand_rt
    ws1[f"D{r}"].alignment = Alignment(horizontal="right", vertical="center")

    ws1.merge_cells(f"D{r+2}:F{r+3}")
    ws1[f"D{r+2}"] = "Affordable Luxury Interiors\nwww.decor8india.com"
    ws1[f"D{r+2}"].font = font_tagline_rt
    ws1[f"D{r+2}"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    r += 5

    # Decorative Divider Line
    ws1.row_dimensions[r].height = 2
    apply_style(ws1, 1, r, 6, r, fill=fill_gold_bar)
    r += 1

    # 4. Client & Proposal Metadata Card (Balanced across A-C and D-F)
    ws1.row_dimensions[r].height = 20
    ws1.merge_cells(f"A{r}:C{r}")
    ws1[f"A{r}"] = "QUOTATION FOR — MR. GOVINDA (Kolte Patil D1304)"
    apply_style(ws1, 1, r, 3, r, font=font_client_hdr, fill=fill_subtotal, alignment=align_left)

    ws1.merge_cells(f"D{r}:F{r}")
    ws1[f"D{r}"] = "Representative: Mr. Satish Bhat (CEO)"
    apply_style(ws1, 4, r, 6, r, font=font_client_hdr, fill=fill_subtotal, alignment=align_right)
    r += 1

    # Symmetrical metadata rows: Col A (lbl1) + Col B:C (val1) | Col D:E (lbl2) + Col F (val2)
    meta_rows = [
        ("Quotation No:", "D8202607102", "Quotation Date:", "06-AUG-2026"),
        ("Contact No:", "+91 93805 23743 / +91 88841 31414", "Project Location:", "Kolte Patil D1304, Bengaluru"),
        ("Project Type:", "Complete 3BHK Turnkey Interior", "Quotation Status:", "Official Luxury Proposal")
    ]
    for m_idx, (lbl1, val1, lbl2, val2) in enumerate(meta_rows):
        ws1.row_dimensions[r].height = 18 if m_idx == 0 else 22

        # Col A: Left Label (width 20 -> fits with no wrapping!)
        ws1[f"A{r}"] = lbl1
        ws1[f"A{r}"].font = font_cell_bold
        ws1[f"A{r}"].fill = fill_meta_lbl
        ws1[f"A{r}"].border = cell_border
        ws1[f"A{r}"].alignment = align_left

        # Col B-C: Left Value (width 36+7 = 43)
        ws1.merge_cells(f"B{r}:C{r}")
        ws1[f"B{r}"] = val1
        apply_style(ws1, 2, r, 3, r, font=font_cell, fill=fill_white, border=cell_border, alignment=align_left)

        # Col D-E: Right Label (width 7+12 = 19 -> fits with no wrapping!)
        ws1.merge_cells(f"D{r}:E{r}")
        ws1[f"D{r}"] = lbl2
        apply_style(ws1, 4, r, 5, r, font=font_cell_bold, fill=fill_meta_lbl, border=cell_border, alignment=align_left)

        # Col F: Right Value (width 15)
        ws1[f"F{r}"] = val2
        ws1[f"F{r}"].font = font_cell_bold if "Official" in val2 else font_cell
        ws1[f"F{r}"].fill = fill_white
        ws1[f"F{r}"].border = cell_border
        ws1[f"F{r}"].alignment = align_right if "2026" in val2 else Alignment(horizontal="left", vertical="center", wrap_text=True)
        r += 1

    # Spacing
    section_subtotal_cells = []
    page_break_rows = []

    # 5. Project Scope & Category Directory (Page 1 Executive Summary)
    ws1.row_dimensions[r].height = 10
    r += 1

    ws1.row_dimensions[r].height = 20
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "PROJECT SCOPE & CATEGORY DIRECTORY"
    apply_style(ws1, 1, r, 6, r, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE), fill=fill_tbl_header, border=header_border, alignment=align_center)
    r += 1

    ws1.row_dimensions[r].height = 18
    ws1[f"A{r}"] = "#"
    ws1[f"A{r}"].font = font_tbl_header
    ws1[f"A{r}"].fill = fill_tbl_header
    ws1[f"A{r}"].border = header_border
    ws1[f"A{r}"].alignment = align_center

    ws1[f"B{r}"] = "Scope Category / Room"
    ws1[f"B{r}"].font = font_tbl_header
    ws1[f"B{r}"].fill = fill_tbl_header
    ws1[f"B{r}"].border = header_border
    ws1[f"B{r}"].alignment = align_left

    ws1.merge_cells(f"C{r}:D{r}")
    ws1[f"C{r}"] = "Key Materials & Design Scope"
    apply_style(ws1, 3, r, 4, r, font=font_tbl_header, fill=fill_tbl_header, border=header_border, alignment=align_left)

    ws1.merge_cells(f"E{r}:F{r}")
    ws1[f"E{r}"] = "Section Sub-Total (₹)"
    apply_style(ws1, 5, r, 6, r, font=font_tbl_header, fill=fill_tbl_header, border=header_border, alignment=align_right)
    r += 1

    scope_directory = [
        ("01", "Modular Kitchen & Sink", "AICA Acrylic + CenturyPly 710 BWP + Carysil Sink"),
        ("02", "Living Room", "Vitrified Marble Wall + Fluted Rafters + 5-Seater Sofa"),
        ("03", "Dining Area", "6-Seater Marble Dining Table & Upholstered Chairs"),
        ("04", "Master Bedroom", "Lacquered Glass Wardrobe + Queen Hydraulic Bed"),
        ("05", "Bedroom 02", "High Gloss Acrylic Wardrobe + King Hydraulic Bed"),
        ("06", "Bedroom 03", "High Gloss Acrylic Wardrobe + King Hydraulic Bed"),
        ("07-09", "Maid, Vanity & Utility", "Carcass 710 BWP + KalingaStone + Jaquar + Storage"),
        ("10-11", "False Ceiling & Electrical", "Gyproc False Ceiling + Havells LED + Chandelier"),
        ("12-14", "Plumbing, Painting & Cleaning", "Asian Paints Royale + Jaquar Plumbing + Deep Clean"),
        ("15-17", "Design, Logistics & Management", "2D/3D Architecture + Site Oversight + Transport"),
    ]
    dir_rows = []
    for s_idx, (s_num, s_name, s_spec) in enumerate(scope_directory):
        ws1.row_dimensions[r].height = 18
        is_zebra = (s_idx % 2 == 1)
        fill_curr = fill_zebra if is_zebra else fill_white

        ws1[f"A{r}"] = s_num
        ws1[f"A{r}"].font = font_cell_bold
        ws1[f"A{r}"].fill = fill_curr
        ws1[f"A{r}"].border = cell_border
        ws1[f"A{r}"].alignment = align_center

        ws1[f"B{r}"] = s_name
        ws1[f"B{r}"].font = font_cell_bold
        ws1[f"B{r}"].fill = fill_curr
        ws1[f"B{r}"].border = cell_border
        ws1[f"B{r}"].alignment = align_left

        ws1.merge_cells(f"C{r}:D{r}")
        ws1[f"C{r}"] = s_spec
        apply_style(ws1, 3, r, 4, r, font=font_cell, fill=fill_curr, border=cell_border, alignment=align_left)

        ws1.merge_cells(f"E{r}:F{r}")
        dir_rows.append(r)
        apply_style(ws1, 5, r, 6, r, font=font_cell_bold, fill=fill_curr, border=cell_border, alignment=align_right)
        r += 1

    # Total Row on Page 1
    dir_total_row = r
    ws1.row_dimensions[r].height = 22
    ws1.merge_cells(f"A{r}:D{r}")
    ws1[f"A{r}"] = "TOTAL ESTIMATED TURNKEY VALUE (INCL. 18% GST)"
    apply_style(ws1, 1, r, 4, r, font=Font(name="Segoe UI", size=10, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)

    ws1.merge_cells(f"E{r}:F{r}")
    apply_style(ws1, 5, r, 6, r, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)
    ws1[f"E{r}"].number_format = FMT_CURRENCY
    r += 1

    # Footnote on Page 1
    ws1.row_dimensions[r].height = 18
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "* Please turn to individual pages ahead for full-scale visual elevations, itemized material specifications, and milestone schedules."
    ws1[f"A{r}"].font = Font(name="Segoe UI", size=8.5, italic=True, color=COLOR_MUTED)
    ws1[f"A{r}"].alignment = align_left
    r += 1

    # PAGE BREAK BEFORE SECTION 01 (Page 1 Cover ends)
    page_break_rows.append(r)

    def calc_row_height(item, spec):
        len_item = len(str(item)) if item else 0
        len_spec = len(str(spec)) if spec else 0
        # Col A (width 20) fits ~19 chars
        # Col B (width 36) fits ~35 chars
        lines_a = (len_item + 18) // 19 if len_item else 1
        lines_b = (len_spec + 34) // 35 if len_spec else 1
        if "\n" in str(spec):
            lines_b = max(lines_b, len(str(spec).split("\n")))
        lines = max(lines_a, lines_b, 1)
        if lines == 1:
            return 18
        elif lines == 2:
            return 23
        elif lines == 3:
            return 32
        else:
            return lines * 11

    def write_section_header(sec_badge, sec_title, sec_sub):
        nonlocal r
        # Badge
        ws1.row_dimensions[r].height = 13
        ws1.merge_cells(f"A{r}:C{r}")
        ws1[f"A{r}"] = sec_badge.upper()
        ws1[f"A{r}"].font = font_sec_badge
        ws1[f"A{r}"].alignment = align_left
        r += 1

        # Title and Top-Right Total Cell
        ws1.row_dimensions[r].height = 20
        ws1.merge_cells(f"A{r}:D{r}")
        ws1[f"A{r}"] = sec_title
        ws1[f"A{r}"].font = font_sec_title
        ws1[f"A{r}"].alignment = align_left

        ws1.merge_cells(f"E{r}:F{r}")
        header_total_cell = f"E{r}"
        ws1[header_total_cell].font = font_sec_total
        ws1[header_total_cell].number_format = FMT_CURRENCY
        ws1[header_total_cell].alignment = align_right
        r += 1

        # Subtitle
        if sec_sub:
            ws1.row_dimensions[r].height = 13
            ws1.merge_cells(f"A{r}:F{r}")
            ws1[f"A{r}"] = sec_sub
            ws1[f"A{r}"].font = font_sec_sub
            ws1[f"A{r}"].alignment = align_left
            r += 1

        # Gold divider
        ws1.row_dimensions[r].height = 2
        apply_style(ws1, 1, r, 6, r, fill=fill_gold_bar)
        r += 1

        return header_total_cell

    def write_table_header(headers=["Item", "Specification & Materials", "Qty", "Unit", "Rate (₹)", "Amount (₹)"]):
        nonlocal r
        ws1.row_dimensions[r].height = 18
        for col_idx, h_text in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            ws1[f"{col_letter}{r}"] = h_text
            ws1[f"{col_letter}{r}"].font = font_tbl_header
            ws1[f"{col_letter}{r}"].fill = fill_tbl_header
            ws1[f"{col_letter}{r}"].border = header_border
            ws1[f"{col_letter}{r}"].alignment = align_center if col_idx in [3, 4] else (align_right if col_idx in [5, 6] else align_left)
        r += 1

    def write_item_row(item, spec, qty, unit, rate, amount_val=None, is_zebra=False):
        nonlocal r
        ws1.row_dimensions[r].height = calc_row_height(item, spec)
        fill_curr = fill_zebra if is_zebra else fill_white

        # Col A: Item
        ws1[f"A{r}"] = item
        ws1[f"A{r}"].font = font_cell_bold
        ws1[f"A{r}"].fill = fill_curr
        ws1[f"A{r}"].border = cell_border
        ws1[f"A{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)

        # Col B: Specification
        ws1[f"B{r}"] = spec
        ws1[f"B{r}"].font = font_cell
        ws1[f"B{r}"].fill = fill_curr
        ws1[f"B{r}"].border = cell_border
        ws1[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)

        # Col C: Qty
        ws1[f"C{r}"] = qty
        ws1[f"C{r}"].font = font_cell
        ws1[f"C{r}"].fill = fill_curr
        ws1[f"C{r}"].border = cell_border
        ws1[f"C{r}"].alignment = align_center

        # Col D: Unit
        ws1[f"D{r}"] = unit
        ws1[f"D{r}"].font = font_cell
        ws1[f"D{r}"].fill = fill_curr
        ws1[f"D{r}"].border = cell_border
        ws1[f"D{r}"].alignment = align_center

        # Col E: Rate
        ws1[f"E{r}"] = rate
        ws1[f"E{r}"].font = font_cell
        ws1[f"E{r}"].fill = fill_curr
        ws1[f"E{r}"].border = cell_border
        if isinstance(rate, (int, float)):
            ws1[f"E{r}"].number_format = FMT_CURRENCY
        ws1[f"E{r}"].alignment = align_right

        # Col F: Amount
        if amount_val is not None:
            ws1[f"F{r}"] = amount_val
        elif isinstance(qty, (int, float)) and isinstance(rate, (int, float)):
            ws1[f"F{r}"] = f"=C{r}*E{r}"
        else:
            ws1[f"F{r}"] = rate
        ws1[f"F{r}"].font = font_cell_amount
        ws1[f"F{r}"].fill = fill_curr
        ws1[f"F{r}"].border = cell_border
        ws1[f"F{r}"].number_format = FMT_CURRENCY
        ws1[f"F{r}"].alignment = align_right

        r += 1

    def write_subtotal_row(label, start_row, spacing_after=4):
        nonlocal r
        subtot_r = r
        ws1.row_dimensions[subtot_r].height = 19
        ws1.merge_cells(f"A{subtot_r}:E{subtot_r}")
        ws1[f"A{subtot_r}"] = label
        apply_style(ws1, 1, subtot_r, 5, subtot_r, font=font_cell_subtot, fill=fill_subtotal, border=subtotal_border, alignment=align_right)

        ws1[f"F{subtot_r}"] = f"=SUM(F{start_row}:F{subtot_r - 1})"
        ws1[f"F{subtot_r}"].font = font_cell_subtot
        ws1[f"F{subtot_r}"].fill = fill_subtotal
        ws1[f"F{subtot_r}"].border = subtotal_border
        ws1[f"F{subtot_r}"].number_format = FMT_CURRENCY
        ws1[f"F{subtot_r}"].alignment = align_right

        ret_cell = f"F{subtot_r}"
        r += 1

        # Spacing
        if spacing_after > 0:
            ws1.row_dimensions[r].height = spacing_after
            r += 1

        return ret_cell

    # =========================================================================
    # --- SECTION 01: MODULAR KITCHEN & SINK PLUMBING ---
    # =========================================================================
    hdr_sec1 = write_section_header("SECTION 01 OF 17", "Modular Kitchen", "AICA High Gloss Acrylic • CenturyPly • Hettich")

    # Left: Large Kitchen Photo placed at A{r}
    if os.path.exists("quatation_renders/kitchen.jpg"):
        img_k = OpenpyxlImage("quatation_renders/kitchen.jpg")
        img_k.width = 240
        img_k.height = 205
        ws1.add_image(img_k, f"A{r}")

    # Right: Specs placed in Columns D to F (merged D-F, NO EMPTY GAP!)
    specs_k = [
        ("Kitchen Finish (Reference Image)", True),
        ("— AICA High Gloss Acrylic Shutters", False),
        ("— CenturyPly 710 BWP — Base Units", False),
        ("— CenturyPly 310 MR — Wall & Loft", False),
        ("— Hettich Hinges & Tandem Channels", False),
        ("— LED Under-Cabinet Lighting", False),
        ("— Factory Manufactured • CNC Cut", False),
    ]
    for idx, (txt, is_b) in enumerate(specs_k):
        curr_r = r + idx
        ws1.row_dimensions[curr_r].height = 18
        ws1.merge_cells(f"D{curr_r}:F{curr_r}")
        ws1[f"D{curr_r}"] = txt
        ws1[f"D{curr_r}"].font = font_bullet_hdr if is_b else font_bullet
        ws1[f"D{curr_r}"].alignment = align_left

    for idx in range(len(specs_k), 12):
        ws1.row_dimensions[r + idx].height = 18

    r += 12

    # Kitchen Units Table
    write_table_header()
    k_start = r
    k_items = [
        ("Base Units", "72 Sq.ft × ₹2,300 — CenturyPly 710 BWP, Acrylic", 72, "Sq.ft", 2300),
        ("Wall Units", "48 Sq.ft × ₹2,150 — CenturyPly 310 MR, Acrylic", 48, "Sq.ft", 2150),
        ("Loft", "48 Sq.ft × ₹1,450 — CenturyPly 310 MR, Laminate", 48, "Sq.ft", 1450),
        ("Kitchen Accessories", "Hettich Cutlery + Pots & Pan + Baskets + Oil Pull-Out + Wicker", 1, "Set", 42000),
        ("Quartz Countertop", "54 Sq.ft × ₹2,200 — KalingaStone Premium", 54, "Sq.ft", 2200),
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(k_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    k_sub_cell = write_subtotal_row("Kitchen Units Sub-Total", start_row=k_start)
    section_subtotal_cells.append(("01 Modular Kitchen (Units)", k_sub_cell))

    # Kitchen Sink & Plumbing Sub-Table
    ws1.row_dimensions[r].height = 16
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "Kitchen Sink & Plumbing"
    ws1[f"A{r}"].font = font_cell_bold
    ws1[f"A{r}"].alignment = align_left
    r += 1

    write_table_header(["Item", "Specification & Materials", "Qty", "Unit", "Rate (₹)", "Amount (₹)"])
    sink_start = r
    sink_items = [
        ("Carysil Premium Quartz Sink", "Composite Quartz Single Bowl with Drainboard", 1, "No.", 18500),
        ("Jaquar Pull-Out Mixer", "Jaquar Chrome Finish Pull-Out Kitchen Faucet", 1, "No.", 12500),
        ("Waste Coupling + Bottle Trap", "Brass & Stainless Steel Heavy Duty Fittings", 2, "Nos", 1750, 3500),
        ("Angle Valves + Flexible Hoses", "Jaquar Certified Chrome Angle Valves & Braided Hoses", 4, "Nos", 725, 2900),
        ("CPVC/UPVC Pipe Materials", "Ashirvad/Supreme Certified Plumbing Line & Couplers", 1, "L.S.", 3500),
        ("Plumbing Consumables", "Solvents, teflon tapes, fasteners, clamps & sealants", 1, "L.S.", 2500),
        ("Kitchen Plumbing Labour", "Installation, alignment, pressure & leakage testing", 1, "L.S.", 5000),
    ]
    for idx, item_data in enumerate(sink_items):
        item, spec, qty, unit, rate = item_data[:5]
        amt = item_data[5] if len(item_data) > 5 else None
        write_item_row(item, spec, qty, unit, rate, amount_val=amt, is_zebra=(idx%2==1))
    sink_sub_cell = write_subtotal_row("Sink & Plumbing Sub-Total", start_row=sink_start)
    section_subtotal_cells.append(("   Kitchen Sink & Plumbing Materials", sink_sub_cell))

    # Update Section 01 Header Total
    ws1[hdr_sec1] = f"={k_sub_cell}+{sink_sub_cell}"

    # Page Break Before Section 02
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 02: LIVING ROOM ---
    # =========================================================================
    hdr_sec2 = write_section_header("SECTION 02 OF 17", "Living Room", "TV Unit • Back Panelling • Feature Wall • Decorative Rafters • Sofa")

    # Balanced 50/50: TV Wall (A-C) and Sofa (D-F)
    if os.path.exists("quatation_renders/living_tv.jpg"):
        img_tv = OpenpyxlImage("quatation_renders/living_tv.jpg")
        img_tv.width = 220
        img_tv.height = 280
        ws1.add_image(img_tv, f"A{r}")

    if os.path.exists("quatation_renders/living_sofa.jpg"):
        img_sofa = OpenpyxlImage("quatation_renders/living_sofa.jpg")
        img_sofa.width = 220
        img_sofa.height = 190
        ws1.add_image(img_sofa, f"D{r}")

    for idx in range(16):
        ws1.row_dimensions[r+idx].height = 18
    r += 16

    living_bullets = [
        ("TV Wall Reference", "Sofa Reference"),
        ("— Marble Finish Feature Wall — Vitrified Tile", "— 5 Seater — Solid Wood Frame"),
        ("— Decorative Wooden Rafters flanking wall", "— High Resilience (HR) Foam"),
        ("— Floating TV Base Unit with LED cove", "— Premium Fabric Upholstery"),
        ("— Back Panelling — CenturyPly 710 BWP", "— Colour: As per design selection"),
        ("— Factory Manufactured & CNC Cut", "")
    ]
    for b_tv, b_sf in living_bullets:
        ws1.row_dimensions[r].height = 16
        ws1.merge_cells(f"A{r}:C{r}")
        ws1[f"A{r}"] = b_tv
        ws1[f"A{r}"].font = font_bullet_hdr if "Reference" in b_tv else font_bullet
        ws1[f"A{r}"].alignment = align_left

        if b_sf:
            ws1.merge_cells(f"D{r}:F{r}")
            ws1[f"D{r}"] = b_sf
            ws1[f"D{r}"].font = font_bullet_hdr if "Reference" in b_sf else font_bullet
            ws1[f"D{r}"].alignment = align_left
        r += 1
    r += 1

    write_table_header()
    lr_start = r
    lr_items = [
        ("TV Unit Base Storage", "16 Sq.ft × ₹1,850 — CenturyPly 710 BWP", 16, "Sq.ft", 1850),
        ("TV Back Panelling", "56 Sq.ft × ₹1,850 — CenturyPly 710 BWP", 56, "Sq.ft", 1850),
        ("Marble Finish Feature Wall", "45 Sq.ft × ₹1,870 — Vitrified Marble Tile", 45, "Sq.ft", 1870),
        ("Decorative Rafters", "54 Sq.ft × ₹1,300 — Engineered Wood", 54, "Sq.ft", 1300),
        ("Premium Sofa", "5 Seater — Solid Wood Frame, HR Foam, Premium Fabric", 1, "Unit", 90000)
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(lr_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    lr_sub_cell = write_subtotal_row("Living Room Sub-Total", start_row=lr_start)
    section_subtotal_cells.append(("02 Living Room", lr_sub_cell))
    ws1[hdr_sec2] = f"={lr_sub_cell}"

    # Page Break Before Section 03 Dining Area
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 03: DINING AREA ---
    # =========================================================================
    hdr_sec3 = write_section_header("SECTION 03 OF 17", "Dining Area", "6-Seater Marble Dining Table & Chairs")

    if os.path.exists("quatation_renders/dining.jpg"):
        img_din = OpenpyxlImage("quatation_renders/dining.jpg")
        img_din.width = 250
        img_din.height = 300
        ws1.add_image(img_din, f"A{r}")

    specs_din = [
        ("Dining Table Reference", True),
        ("— 6-Seater Marble-Top Dining Table", False),
        ("— Fluted Column Pedestal Base", False),
        ("— Gold Accent Edge Detail", False),
        ("— 6 Upholstered Dining Chairs included", False),
        ("— Premium Fabric — Cream Tone", False),
    ]
    for idx, (txt, is_b) in enumerate(specs_din):
        curr_r = r + idx
        ws1.row_dimensions[curr_r].height = 20
        ws1.merge_cells(f"D{curr_r}:F{curr_r}")
        ws1[f"D{curr_r}"] = txt
        ws1[f"D{curr_r}"].font = font_bullet_hdr if is_b else font_bullet
        ws1[f"D{curr_r}"].alignment = align_left

    for idx in range(len(specs_din), 17):
        ws1.row_dimensions[r + idx].height = 18

    r += 17

    write_table_header()
    din_start = r
    write_item_row("Dining Table Set", "6-Seater Marble-Top Dining Table + 6 Upholstered Chairs", 1, "Set", 142700)
    din_sub_cell = write_subtotal_row("Dining Area Sub-Total", start_row=din_start)
    section_subtotal_cells.append(("03 Dining Area", din_sub_cell))
    ws1[hdr_sec3] = f"={din_sub_cell}"

    # Page Break Before Section 04
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 04: MASTER BEDROOM ---
    # =========================================================================
    hdr_sec4 = write_section_header("SECTION 04 OF 17", "Master Bedroom", "Lacquered Glass Wardrobe • Queen Bed with Hydraulic Storage • Side Tables")

    if os.path.exists("quatation_renders/mbr_wardrobe.jpg"):
        img_mbr_w = OpenpyxlImage("quatation_renders/mbr_wardrobe.jpg")
        img_mbr_w.width = 220
        img_mbr_w.height = 280
        ws1.add_image(img_mbr_w, f"A{r}")

    if os.path.exists("quatation_renders/mbr_bed.jpg"):
        img_mbr_b = OpenpyxlImage("quatation_renders/mbr_bed.jpg")
        img_mbr_b.width = 220
        img_mbr_b.height = 300
        ws1.add_image(img_mbr_b, f"D{r}")

    for idx in range(17):
        ws1.row_dimensions[r+idx].height = 18
    r += 17

    mbr_bullets = [
        ("Master Bedroom Wardrobe", "Bed Reference"),
        ("— Premium Lacquered Glass Sliding Shutters", "— Queen Size Bed — Hydraulic Storage"),
        ("— CenturyPly 310 MR — 19 mm Carcass, 9 mm Back", "— Premium Upholstered Headboard"),
        ("— Hettich Premium Sliding Hardware & Soft-Close", "— Two Bedside Tables with storage"),
        ("— 4 Soft-Close Drawers • 2 Hanging Sections", "— Soft-Close Hardware throughout"),
        ("— 8 Adjustable Shelves • Loft Included", "— Dressing Area with backlit round mirror"),
        ("— CNC Machine Manufactured", "")
    ]
    for b_w, b_b in mbr_bullets:
        ws1.row_dimensions[r].height = 16
        ws1.merge_cells(f"A{r}:C{r}")
        ws1[f"A{r}"] = b_w
        ws1[f"A{r}"].font = font_bullet_hdr if "Wardrobe" in b_w else font_bullet
        ws1[f"A{r}"].alignment = align_left

        if b_b:
            ws1.merge_cells(f"D{r}:F{r}")
            ws1[f"D{r}"] = b_b
            ws1[f"D{r}"].font = font_bullet_hdr if "Bed" in b_b else font_bullet
            ws1[f"D{r}"].alignment = align_left
        r += 1
    r += 1

    write_table_header()
    mbr_start = r
    mbr_items = [
        ("Wardrobe with Loft", "108 Sq.ft × ₹2,350 — Lacquered Glass Sliding Shutters", 108, "Sq.ft", 2350),
        ("Queen Size Bed", "Hydraulic Storage + Upholstered Headboard + 2 Side Tables", 1, "Unit", 127800)
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(mbr_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    mbr_sub_cell = write_subtotal_row("Master Bedroom Sub-Total", start_row=mbr_start)
    section_subtotal_cells.append(("04 Master Bedroom", mbr_sub_cell))
    ws1[hdr_sec4] = f"={mbr_sub_cell}"

    # Page Break Before Section 05 Bedroom 02
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 05: BEDROOM 02 ---
    # =========================================================================
    hdr_sec5 = write_section_header("SECTION 05 OF 17", "Bedroom 02", "Acrylic Sliding Wardrobe • Loft • King Bed with Hydraulic Storage")

    if os.path.exists("quatation_renders/bed2_wardrobe.jpg"):
        img_b2 = OpenpyxlImage("quatation_renders/bed2_wardrobe.jpg")
        img_b2.width = 240
        img_b2.height = 340
        ws1.add_image(img_b2, f"A{r}")

    specs_b2 = [
        ("Bedroom 02 Wardrobe Reference", True),
        ("— High Gloss Acrylic Sliding Shutters", False),
        ("— CenturyPly 310 MR — 19mm Carcass", False),
        ("— Hettich Sliding Hardware", False),
        ("— 3 Soft-Close Drawers Included", False),
        ("— Hanging Section & Adjustable Shelves", False),
    ]
    for idx, (txt, is_b) in enumerate(specs_b2):
        curr_r = r + idx
        ws1.row_dimensions[curr_r].height = 20
        ws1.merge_cells(f"D{curr_r}:F{curr_r}")
        ws1[f"D{curr_r}"] = txt
        ws1[f"D{curr_r}"].font = font_bullet_hdr if is_b else font_bullet
        ws1[f"D{curr_r}"].alignment = align_left

    for idx in range(len(specs_b2), 19):
        ws1.row_dimensions[r + idx].height = 18

    r += 19

    write_table_header()
    b2_start = r
    b2_items = [
        ("Wardrobe", "77 Sq.ft × ₹2,100 — Acrylic Sliding Shutters", 77, "Sq.ft", 2100),
        ("Loft", "22 Sq.ft × ₹1,450 — CenturyPly 310 MR, Laminate", 22, "Sq.ft", 1450),
        ("King Size Bed", "Hydraulic Storage • Premium Headboard • Factory Finished", 1, "Unit", 120388)
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(b2_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    b2_sub_cell = write_subtotal_row("Bedroom 02 Sub-Total", start_row=b2_start)
    section_subtotal_cells.append(("05 Bedroom 02", b2_sub_cell))
    ws1[hdr_sec5] = f"={b2_sub_cell}"

    # Page Break Before Section 06
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 06: BEDROOM 03 ---
    # =========================================================================
    hdr_sec6 = write_section_header("SECTION 06 OF 17", "Bedroom 03", "Acrylic Sliding Wardrobe • Loft • King Bed with Hydraulic Storage")

    if os.path.exists("quatation_renders/bed3_wardrobe.jpg"):
        img_b3 = OpenpyxlImage("quatation_renders/bed3_wardrobe.jpg")
        img_b3.width = 240
        img_b3.height = 305
        ws1.add_image(img_b3, f"A{r}")

    specs_b3 = [
        ("Bedroom 03 Wardrobe Reference", True),
        ("— High Gloss Acrylic Sliding Shutters", False),
        ("— CenturyPly 310 MR — 19 mm Carcass", False),
        ("— Hettich Sliding Hardware", False),
        ("— 3 Soft-Close Drawers Included", False),
        ("— Hanging Section & Adjustable Shelves", False),
        ("— Factory Manufactured • CNC Cut", False),
    ]
    for idx, (txt, is_b) in enumerate(specs_b3):
        curr_r = r + idx
        ws1.row_dimensions[curr_r].height = 20
        ws1.merge_cells(f"D{curr_r}:F{curr_r}")
        ws1[f"D{curr_r}"] = txt
        ws1[f"D{curr_r}"].font = font_bullet_hdr if is_b else font_bullet
        ws1[f"D{curr_r}"].alignment = align_left

    for idx in range(len(specs_b3), 17):
        ws1.row_dimensions[r + idx].height = 18

    r += 17

    write_table_header()
    b3_start = r
    b3_items = [
        ("Wardrobe", "77 Sq.ft × ₹2,100 — Acrylic Sliding Shutters", 77, "Sq.ft", 2100),
        ("Loft", "22 Sq.ft × ₹1,450 — CenturyPly 310 MR, Laminate", 22, "Sq.ft", 1450),
        ("King Size Bed", "Hydraulic Storage • Premium Headboard • Factory Finished", 1, "Unit", 120300)
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(b3_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    b3_sub_cell = write_subtotal_row("Bedroom 03 Sub-Total", start_row=b3_start)
    section_subtotal_cells.append(("06 Bedroom 03", b3_sub_cell))
    ws1[hdr_sec6] = f"={b3_sub_cell}"

    # Page Break Before Section 07 Maid Room
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 07: MAID ROOM ---
    # =========================================================================
    hdr_sec7 = write_section_header("SECTION 07 OF 17", "Maid Room", "Single Bed with Hydraulic Storage")
    write_table_header()
    m_start = r
    write_item_row("Single Bed", "CenturyPly 310 MR • Hydraulic Storage • Laminate Finish", 1, "Unit", 30525)
    m_sub_cell = write_subtotal_row("Maid Room Sub-Total", start_row=m_start)
    section_subtotal_cells.append(("07 Maid Room", m_sub_cell))
    ws1[hdr_sec7] = f"={m_sub_cell}"

    # =========================================================================
    # --- SECTION 08: BATHROOM VANITY UNITS ---
    # =========================================================================
    hdr_sec8 = write_section_header("SECTION 08 OF 17", "Bathroom Vanity Units", "Countertops, Mirrors & Sanitary Mixers for 3 Bathrooms")

    inc_items = [
        "✓ INCLUDED  KalingaStone Quartz Top — Premium Quartz Countertop — 3 Bathrooms",
        "✓ INCLUDED  Premium Mirror — 3 Nos — Framed/Backlit per design",
        "✓ INCLUDED  Jaquar Basin Mixer — 3 Nos — Premium Collection",
        "✓ INCLUDED  Plumbing Fittings — Bottle Trap, Waste Coupling, Angle Valve, Hoses — 3 sets"
    ]
    for inc_txt in inc_items:
        ws1.row_dimensions[r].height = 18
        ws1.merge_cells(f"A{r}:F{r}")
        ws1[f"A{r}"] = inc_txt
        ws1[f"A{r}"].font = font_included
        ws1[f"A{r}"].fill = fill_zebra
        ws1[f"A{r}"].alignment = align_left
        r += 1
    r += 1

    write_table_header()
    v_start = r
    write_item_row("Vanity Carcass", "29 Sq.ft × ₹2,300 — CenturyPly 710 BWP 19mm Marine Grade", 29, "Sq.ft", 2300)
    v_sub_cell = write_subtotal_row("Bathroom Vanity Sub-Total", start_row=v_start)
    section_subtotal_cells.append(("08 Bathroom Vanity Units", v_sub_cell))
    ws1[hdr_sec8] = f"={v_sub_cell}"

    # =========================================================================
    # --- SECTION 09: UTILITY STORAGE ---
    # =========================================================================
    hdr_sec9 = write_section_header("SECTION 09 OF 17", "Utility Storage", "Base + Wall + Loft Storage Unit")
    write_table_header()
    u_start = r
    write_item_row("Utility Storage Unit", "25 Sq.ft × ₹1,950 — Base + Wall + Loft • CenturyPly 710 BWP • Laminate", 25, "Sq.ft", 1950)
    u_sub_cell = write_subtotal_row("Utility Storage Sub-Total", start_row=u_start)
    section_subtotal_cells.append(("09 Utility Storage", u_sub_cell))
    ws1[hdr_sec9] = f"={u_sub_cell}"

    # Page Break Before Section 10
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 10: FALSE CEILING ---
    # =========================================================================
    hdr_sec10 = write_section_header("SECTION 10 OF 17", "False Ceiling", "Gypsum False Ceiling, PVC System & Grid Ceilings")
    write_table_header()
    fc_start = r
    fc_items = [
        ("Gypsum False Ceiling — All Areas", "Saint-Gobain Gyproc 12mm • GI Channel • Cove + Profile + LED Provision", 1, "L.S.", 219700),
        ("PVC Ceiling — Utility Area", "Waterproof PVC Panel System with perimeter finish", 1, "L.S.", 20000),
        ("Grid Ceiling — Service Area", "Armstrong / Gyproc Grid System with moisture resistance", 1, "L.S.", 20500)
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(fc_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    fc_sub_cell = write_subtotal_row("False Ceiling Sub-Total", start_row=fc_start)
    section_subtotal_cells.append(("10 False Ceiling", fc_sub_cell))
    ws1[hdr_sec10] = f"={fc_sub_cell}"

    # =========================================================================
    # --- SECTION 11: ELECTRICAL & LIGHTING ---
    # =========================================================================
    hdr_sec11 = write_section_header("SECTION 11 OF 17", "Electrical & Lighting", "Havells LED Lights, Strip Lights, Chandelier & Profile Lighting")
    write_table_header()
    el_start = r
    el_items = [
        ("Electrical Labour", "Complete wiring, conduits, modular switch fittings & testing", 1, "L.S.", 15000),
        ("Havells 12W LED Panel Lights", "26 Nos × ₹650 — Warm white / natural white ceiling panel lights", 26, "Nos", 650),
        ("Havells COB Spot Lights", "12 Nos × ₹950 — Focused architectural spotlighting", 12, "Nos", 950),
        ("Havells LED Strip Lights (5m Roll)", "8 Rolls × ₹1,365 — 3000K warm cove accent strip lights", 8, "Rolls", 1365),
        ("Havells LED Drivers", "8 Nos × ₹650 — Constant voltage dedicated driver units", 8, "Nos", 650),
        ("Decorative Pendant Light", "1 No. × ₹6,500 — Dining / Foyer designer accent light", 1, "No.", 6500),
        ("Living Room Chandelier", "1 No. × ₹18,000 — Designer luxury living room centerpiece", 1, "No.", 18000),
        ("Aluminium Profile Lights", "40 Rft × ₹650/Rft — Recessed slim profile with diffuser", 40, "Rft", 650),
        ("FRLS Wires & Modular Switches", "Polycab/Finolex FRLS Wires, Anchor/Schneider modular switches", 1, "L.S.", 20000)
    ]
    for idx, (item, spec, qty, unit, rate) in enumerate(el_items):
        write_item_row(item, spec, qty, unit, rate, is_zebra=(idx%2==1))
    el_sub_cell = write_subtotal_row("Electrical & Lighting Sub-Total", start_row=el_start)
    section_subtotal_cells.append(("11 Electrical & Lighting", el_sub_cell))
    ws1[hdr_sec11] = f"={el_sub_cell}"

    # Page Break Before Section 12 Plumbing Labour
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 12: PLUMBING LABOUR ---
    # =========================================================================
    hdr_sec12 = write_section_header("SECTION 12 OF 17", "Plumbing Labour", "Sanitaryware, Basin Mixers & Drainline Connections")
    write_table_header()
    pl_start = r
    write_item_row("Plumbing Labour", "Kitchen Sink + 3 Basin Mixers + Connections + Leakage Testing", 1, "L.S.", 12000)
    pl_sub_cell = write_subtotal_row("Plumbing Sub-Total", start_row=pl_start)
    section_subtotal_cells.append(("12 Plumbing Labour", pl_sub_cell))
    ws1[hdr_sec12] = f"={pl_sub_cell}"

    # =========================================================================
    # --- SECTION 13: PAINTING ---
    # =========================================================================
    hdr_sec13 = write_section_header("SECTION 13 OF 17", "Painting", "Asian Paints Royale Luxury Emulsion with Full Surface Prep")
    write_table_header()
    pt_start = r
    write_item_row("Interior Painting", "3,850 Sq.ft × ₹42 — 2 coats Wall Putty + 1 Primer + 2 coats Royale Luxury", 3850, "Sq.ft", 42)
    pt_sub_cell = write_subtotal_row("Painting Sub-Total", start_row=pt_start)
    section_subtotal_cells.append(("13 Painting", pt_sub_cell))
    ws1[hdr_sec13] = f"={pt_sub_cell}"

    # =========================================================================
    # --- SECTION 14: DEEP CLEANING ---
    # =========================================================================
    hdr_sec14 = write_section_header("SECTION 14 OF 17", "Deep Cleaning", "Post-Installation Professional House Deep Cleaning")
    write_table_header()
    dc_start = r
    write_item_row("Deep Cleaning", "Post-installation professional deep clean, vacuuming & sanitization", 1, "Job", 15000)
    dc_sub_cell = write_subtotal_row("Deep Cleaning Sub-Total", start_row=dc_start)
    section_subtotal_cells.append(("14 Deep Cleaning", dc_sub_cell))
    ws1[hdr_sec14] = f"={dc_sub_cell}"

    # Page Break Before Section 15 Design Fees
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 15: DESIGN FEES ---
    # =========================================================================
    hdr_sec15 = write_section_header("SECTION 15 OF 17", "Design Fees", "2D/3D Architecture, Working Drawings & Walkthrough Video")

    des_bullets = [
        ("— 2D Layout Planning", "— Plumbing Drawings"),
        ("— 3D Design Visualization", "— False Ceiling Drawings"),
        ("— Working Drawings", "— Walkthrough Video"),
        ("— Electrical Drawings", "— Material Moodboard & Palette")
    ]
    for b_l, b_r in des_bullets:
        ws1.row_dimensions[r].height = 18
        ws1.merge_cells(f"A{r}:C{r}")
        ws1[f"A{r}"] = b_l
        ws1[f"A{r}"].font = font_bullet
        ws1[f"A{r}"].alignment = align_left

        ws1.merge_cells(f"D{r}:F{r}")
        ws1[f"D{r}"] = b_r
        ws1[f"D{r}"].font = font_bullet
        ws1[f"D{r}"].alignment = align_left
        r += 1
    r += 1

    write_table_header()
    df_start = r
    write_item_row("Design Fees", "1,200 Sq.ft × ₹55/Sq.ft — Architectural & 3D Interior Design Package", 1200, "Sq.ft", 55)
    df_sub_cell = write_subtotal_row("Design Fees Sub-Total", start_row=df_start)
    section_subtotal_cells.append(("15 Design Fees", df_sub_cell))
    ws1[hdr_sec15] = f"={df_sub_cell}"

    # =========================================================================
    # --- SECTION 16: LOADING & UNLOADING ---
    # =========================================================================
    hdr_sec16 = write_section_header("SECTION 16 OF 17", "Loading & Unloading", "Site Logistics & Material Transport Coordination")
    write_table_header()
    lu_start = r
    write_item_row("Loading & Unloading", "Site logistics — material handling, transport coordination, vertical shifting", 1, "Job", 12000)
    lu_sub_cell = write_subtotal_row("Loading & Unloading Sub-Total", start_row=lu_start)
    section_subtotal_cells.append(("16 Loading & Unloading", lu_sub_cell))
    ws1[hdr_sec16] = f"={lu_sub_cell}"

    # =========================================================================
    # --- SECTION 17: PROJECT MANAGEMENT ---
    # =========================================================================
    hdr_sec17 = write_section_header("SECTION 17 OF 17", "Project Management", "Site Supervision, Quality Control & Project Handover")

    pm_bullets = [
        ("— Site Supervision", "— Execution Management"),
        ("— Vendor Coordination", "— Final Snag Rectification"),
        ("— Material Inspection", "— Project Handover"),
        ("— Quality Control", "— Client Alignment Meetings")
    ]
    for b_l, b_r in pm_bullets:
        ws1.row_dimensions[r].height = 18
        ws1.merge_cells(f"A{r}:C{r}")
        ws1[f"A{r}"] = b_l
        ws1[f"A{r}"].font = font_bullet
        ws1[f"A{r}"].alignment = align_left

        ws1.merge_cells(f"D{r}:F{r}")
        ws1[f"D{r}"] = b_r
        ws1[f"D{r}"].font = font_bullet
        ws1[f"D{r}"].alignment = align_left
        r += 1
    r += 1

    write_table_header()
    pm_start = r
    write_item_row("Project Management", "Site supervision & execution oversight, full project tenure", 1, "Job", 42000)
    pm_sub_cell = write_subtotal_row("Project Management Sub-Total", start_row=pm_start)
    section_subtotal_cells.append(("17 Project Management", pm_sub_cell))
    ws1[hdr_sec17] = f"={pm_sub_cell}"

    # Page Break Before Section 18 Grand Summary
    page_break_rows.append(r)

    # =========================================================================
    # --- SECTION 18: GRAND SUMMARY (EXACT MATCH TO REFERENCE IMAGE 12) ---
    # =========================================================================
    ws1.row_dimensions[r].height = 4
    apply_style(ws1, 1, r, 6, r, fill=fill_gold_bar)
    r += 1

    ws1.row_dimensions[r].height = 28
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "Grand Summary"
    ws1[f"A{r}"].font = Font(name="Segoe UI", size=15, bold=True, color=COLOR_DARK)
    ws1[f"A{r}"].alignment = align_left
    r += 1

    ws1.row_dimensions[r].height = 16
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "Complete cost breakdown across all 17 sections"
    ws1[f"A{r}"].font = font_sec_sub
    ws1[f"A{r}"].alignment = align_left
    r += 1

    # Table Header across A to F
    ws1.row_dimensions[r].height = 20
    ws1[f"A{r}"] = "#"
    ws1[f"A{r}"].font = font_tbl_header
    ws1[f"A{r}"].fill = fill_tbl_header
    ws1[f"A{r}"].border = header_border
    ws1[f"A{r}"].alignment = align_center

    ws1.merge_cells(f"B{r}:D{r}")
    ws1[f"B{r}"] = "Section"
    apply_style(ws1, 2, r, 4, r, font=font_tbl_header, fill=fill_tbl_header, border=header_border, alignment=align_left)

    ws1.merge_cells(f"E{r}:F{r}")
    ws1[f"E{r}"] = "Amount (₹)"
    apply_style(ws1, 5, r, 6, r, font=font_tbl_header, fill=fill_tbl_header, border=header_border, alignment=align_right)
    r += 1

    summary_start_row = r
    for idx, (sec_name, cell_ref) in enumerate(section_subtotal_cells):
        ws1.row_dimensions[r].height = 20
        is_zebra = (idx % 2 == 1)
        fill_curr = fill_zebra if is_zebra else fill_white

        sec_num = sec_name.split()[0] if sec_name.strip()[0].isdigit() else ""
        ws1[f"A{r}"] = sec_num
        ws1[f"A{r}"].font = font_cell_bold
        ws1[f"A{r}"].fill = fill_curr
        ws1[f"A{r}"].border = cell_border
        ws1[f"A{r}"].alignment = align_center

        sec_clean_title = sec_name[len(sec_num):].strip() if sec_num else sec_name.strip()
        ws1.merge_cells(f"B{r}:D{r}")
        ws1[f"B{r}"] = sec_clean_title
        apply_style(ws1, 2, r, 4, r, font=font_cell, fill=fill_curr, border=cell_border, alignment=align_left)

        ws1.merge_cells(f"E{r}:F{r}")
        ws1[f"E{r}"] = f"={cell_ref}"
        apply_style(ws1, 5, r, 6, r, font=font_cell_amount, fill=fill_curr, border=cell_border, alignment=align_right)
        ws1[f"E{r}"].number_format = FMT_CURRENCY

        r += 1

    summary_end_row = r - 1

    # GRAND TOTAL BANNER
    ws1.row_dimensions[r].height = 28
    ws1.merge_cells(f"A{r}:D{r}")
    ws1[f"A{r}"] = "GRAND TOTAL"
    apply_style(ws1, 1, r, 4, r, font=font_grand_lbl, fill=fill_grand_total, border=grand_total_border, alignment=align_left)

    grand_total_cell = f"E{r}"
    ws1.merge_cells(f"E{r}:F{r}")
    ws1[f"E{r}"] = f"=SUM(E{summary_start_row}:E{summary_end_row})"
    apply_style(ws1, 5, r, 6, r, font=font_grand_val, fill=fill_grand_total, border=grand_total_border, alignment=align_right)
    ws1[f"E{r}"].number_format = FMT_CURRENCY
    r += 1

    # Computed note
    ws1.row_dimensions[r].height = 18
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "* Computed total: ₹29,22,133. GST @18% extra. Final billing per site measurement."
    ws1[f"A{r}"].font = Font(name="Segoe UI", size=8.5, italic=True, color=COLOR_MUTED)
    ws1[f"A{r}"].alignment = align_left
    r += 1

    # TAX BREAKDOWN (CGST 9% + SGST 9%)
    tax_rows = [
        ("CGST @ 9%", f"=ROUND({grand_total_cell}*0.09, 0)"),
        ("SGST @ 9%", f"=ROUND({grand_total_cell}*0.09, 0)"),
    ]
    cgst_cell, sgst_cell = None, None
    for idx, (tax_lbl, tax_formula) in enumerate(tax_rows):
        ws1.row_dimensions[r].height = 20
        ws1.merge_cells(f"A{r}:D{r}")
        ws1[f"A{r}"] = tax_lbl
        apply_style(ws1, 1, r, 4, r, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)

        ws1.merge_cells(f"E{r}:F{r}")
        ws1[f"E{r}"] = tax_formula
        apply_style(ws1, 5, r, 6, r, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
        ws1[f"E{r}"].number_format = FMT_CURRENCY
        if idx == 0: cgst_cell = f"E{r}"
        else: sgst_cell = f"E{r}"
        r += 1

    # TOTAL PROJECT COST (INCL. GST)
    ws1.row_dimensions[r].height = 24
    ws1.merge_cells(f"A{r}:D{r}")
    ws1[f"A{r}"] = "TOTAL PROJECT COST (A+B+C) INCLUDING 18% GST"
    apply_style(ws1, 1, r, 4, r, font=Font(name="Segoe UI", size=10, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)

    total_incl_gst_cell = f"E{r}"
    ws1.merge_cells(f"E{r}:F{r}")
    ws1[f"E{r}"] = f"={grand_total_cell}+{cgst_cell}+{sgst_cell}"
    apply_style(ws1, 5, r, 6, r, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)
    ws1[f"E{r}"].number_format = FMT_CURRENCY
    r += 2

    # PAYMENT SCHEDULE
    ws1.row_dimensions[r].height = 20
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "Payment Schedule"
    apply_style(ws1, 1, r, 6, r, font=font_sec_title, fill=fill_subtotal, border=header_border, alignment=align_left)
    r += 1

    ws1.row_dimensions[r].height = 19
    ws1[f"A{r}"] = "Sl No"
    ws1[f"A{r}"].font = font_tbl_header
    ws1[f"A{r}"].fill = fill_tbl_header
    ws1[f"A{r}"].border = header_border
    ws1[f"A{r}"].alignment = align_center

    ws1.merge_cells(f"B{r}:C{r}")
    ws1[f"B{r}"] = "Payment Timelines & Milestones"
    apply_style(ws1, 2, r, 3, r, font=font_tbl_header, fill=fill_tbl_header, border=header_border, alignment=align_left)

    ws1[f"D{r}"] = "Stage %"
    ws1[f"D{r}"].font = font_tbl_header
    ws1[f"D{r}"].fill = fill_tbl_header
    ws1[f"D{r}"].border = header_border
    ws1[f"D{r}"].alignment = align_center

    ws1.merge_cells(f"E{r}:F{r}")
    ws1[f"E{r}"] = "Milestone Amount (₹)"
    apply_style(ws1, 5, r, 6, r, font=font_tbl_header, fill=fill_tbl_header, border=header_border, alignment=align_right)
    r += 1

    payment_milestones = [
        (1, "At the time of project / order confirmation", 0.15),
        (2, "At the time of starting the production / execution at factory", 0.50),
        (3, "After completion of production / before dispatch to site for installation", 0.35)
    ]
    pay_start_r = r
    for idx_m, (sl, timeline, pct) in enumerate(payment_milestones):
        ws1.row_dimensions[r].height = 22
        ws1[f"A{r}"] = sl
        ws1[f"A{r}"].font = font_cell_bold
        ws1[f"A{r}"].fill = fill_white
        ws1[f"A{r}"].border = cell_border
        ws1[f"A{r}"].alignment = align_center

        ws1.merge_cells(f"B{r}:C{r}")
        ws1[f"B{r}"] = timeline
        apply_style(ws1, 2, r, 3, r, font=font_cell, fill=fill_white, border=cell_border, alignment=align_left)

        ws1[f"D{r}"] = pct
        ws1[f"D{r}"].font = font_cell_bold
        ws1[f"D{r}"].fill = fill_white
        ws1[f"D{r}"].border = cell_border
        ws1[f"D{r}"].number_format = '0%'
        ws1[f"D{r}"].alignment = align_center

        ws1.merge_cells(f"E{r}:F{r}")
        if idx_m == 2:
            ws1[f"E{r}"] = f"={total_incl_gst_cell}-E{pay_start_r}-E{pay_start_r+1}"
        else:
            ws1[f"E{r}"] = f"=ROUND({total_incl_gst_cell}*D{r}, 0)"
        apply_style(ws1, 5, r, 6, r, font=font_cell_bold, fill=fill_white, border=cell_border, alignment=align_right)
        ws1[f"E{r}"].number_format = FMT_CURRENCY
        r += 1

    ws1.row_dimensions[r].height = 20
    ws1.merge_cells(f"A{r}:C{r}")
    ws1[f"A{r}"] = "Total Milestone Payment Check"
    apply_style(ws1, 1, r, 3, r, font=font_cell_subtot, fill=fill_subtotal, border=subtotal_border, alignment=align_right)

    ws1[f"D{r}"] = f"=SUM(D{pay_start_r}:D{r-1})"
    ws1[f"D{r}"].font = font_cell_subtot
    ws1[f"D{r}"].fill = fill_subtotal
    ws1[f"D{r}"].border = subtotal_border
    ws1[f"D{r}"].number_format = '0%'
    ws1[f"D{r}"].alignment = align_center

    ws1.merge_cells(f"E{r}:F{r}")
    ws1[f"E{r}"] = f"=SUM(E{pay_start_r}:E{r-1})"
    apply_style(ws1, 5, r, 6, r, font=font_cell_subtot, fill=fill_subtotal, border=subtotal_border, alignment=align_right)
    ws1[f"E{r}"].number_format = FMT_CURRENCY
    r += 1

    # Page Break Before Company Account Information & Terms
    page_break_rows.append(r)

    # COMPANY ACCOUNT DETAILS
    ws1.row_dimensions[r].height = 22
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "Company Account Information"
    apply_style(ws1, 1, r, 6, r, font=font_sec_title, fill=fill_subtotal, border=header_border, alignment=align_left)
    r += 1

    ws1.row_dimensions[r].height = 18
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = 'Please issue A/c Payee cheque in the name of "DECOR8 INDIA" or transfer online to:'
    ws1[f"A{r}"].font = font_cell_bold
    ws1[f"A{r}"].alignment = align_left
    r += 1

    bank_details = [
        ("Bank Name:", "IDFC FIRST BANK", "Account Name:", "DECOR8 INDIA"),
        ("Branch Name:", "BTM Layout, Bengaluru", "Account Number:", "10075641203"),
        ("Account Type:", "Current Account", "IFSC Code:", "IDFB0080182")
    ]
    for lbl1, val1, lbl2, val2 in bank_details:
        ws1.row_dimensions[r].height = 20
        # Col A: Left Label (width 20)
        ws1[f"A{r}"] = lbl1
        ws1[f"A{r}"].font = font_cell_bold
        ws1[f"A{r}"].fill = fill_meta_lbl
        ws1[f"A{r}"].border = cell_border
        ws1[f"A{r}"].alignment = align_left

        # Col B-C: Left Value (width 36+7 = 43)
        ws1.merge_cells(f"B{r}:C{r}")
        ws1[f"B{r}"] = val1
        apply_style(ws1, 2, r, 3, r, font=font_cell_bold if "IDFC" in val1 else font_cell, fill=fill_white, border=cell_border, alignment=align_left)

        # Col D-E: Right Label (width 7+12 = 19 -> NO WRAPPING!)
        ws1.merge_cells(f"D{r}:E{r}")
        ws1[f"D{r}"] = lbl2
        apply_style(ws1, 4, r, 5, r, font=font_cell_bold, fill=fill_meta_lbl, border=cell_border, alignment=align_left)

        # Col F: Right Value (width 15)
        ws1[f"F{r}"] = val2
        ws1[f"F{r}"].font = font_cell_bold
        ws1[f"F{r}"].fill = fill_white
        ws1[f"F{r}"].border = cell_border
        ws1[f"F{r}"].alignment = align_left
        r += 1
    r += 1

    # TERMS & CONDITIONS
    ws1.row_dimensions[r].height = 22
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "* Terms & Conditions"
    apply_style(ws1, 1, r, 6, r, font=font_terms_hdr, fill=fill_subtotal, border=header_border, alignment=align_left)
    r += 1

    terms = [
        "* Price is valid for 30 days only after the 1st estimation is offered. Possibility of price revision may happen due to vendor supplier reasons.",
        "* DECOR8 will cover warranty period upto 15 years for any wood work related products on manufacturing defects (Plywood bendness not included). DECOR8 will not cover warranty for intentional wear and tear or damages from customer side. All other products from different vendors like Hettich, KAFF, EBCO, Havells etc will be covered by respective vendors as per their norms and DECOR8 will support for the same.",
        "* Payment terms as per payment slab should be followed from the time of booking by the client for smooth functioning of work. After any revised quote due to addition of products/changes, if price adds up, client needs to pay difference amount during that phase. Any reductions will be adjusted in final billing.",
        "* Designs will not start until complete booking amount is credited to the company followed by payment slabs for production & installation as per progress.",
        "* Any delay in terms of payment as per above payment schedule may cause delay in execution for which DECOR8 won't be responsible.",
        "* Price may be revised if there is any official raw material price hike by supplier while project is still in designing stage. No revision for projects already under production.",
        "* Once an order is confirmed and customer wants to cancel the order, NO amount will be refunded.",
        "* No hidden charges for 3D designs (only 03 corrections/revisions allowed).",
        "* Any corrections beyond 03 times, INR 3,000 per each rendering will be charged to the client.",
        "* Once the designs are approved and execution has started, no changes will be accepted in terms of structure or design modifications."
    ]
    for t_txt in terms:
        lines = max((len(t_txt) + 80) // 85, 1)
        ws1.row_dimensions[r].height = max(lines * 13, 18)
        ws1.merge_cells(f"A{r}:F{r}")
        ws1[f"A{r}"] = t_txt
        ws1[f"A{r}"].font = font_terms_item
        ws1[f"A{r}"].alignment = align_left_top
        r += 1
    r += 1

    # Signatures
    align_sig = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws1.row_dimensions[r].height = 42
    ws1.merge_cells(f"A{r}:C{r}")
    ws1[f"A{r}"] = "Authorized Signatory\nFor DECOR8 INDIA (Satish Bhat, CEO)"
    apply_style(ws1, 1, r, 3, r, font=font_cell_bold, fill=fill_white, border=cell_border, alignment=align_sig)

    ws1.merge_cells(f"D{r}:F{r}")
    ws1[f"D{r}"] = "Client Acceptance Signature\nMR. GOVINDA (Kolte Patil D1304)"
    apply_style(ws1, 4, r, 6, r, font=font_cell_bold, fill=fill_white, border=cell_border, alignment=align_sig)
    r += 2

    ws1.row_dimensions[r].height = 20
    ws1.merge_cells(f"A{r}:F{r}")
    ws1[f"A{r}"] = "www.decor8india.com  •  Affordable Luxury Turnkey Interiors  •  Bengaluru"
    apply_style(ws1, 1, r, 6, r, font=Font(name="Segoe UI", size=9, bold=True, color=COLOR_GOLD_DARK), fill=fill_card, alignment=align_center)

    # Link Page 1 Directory Table to actual subtotal cells
    ws1[f"E{dir_rows[0]}"] = f"={k_sub_cell}+{sink_sub_cell}"
    ws1[f"E{dir_rows[1]}"] = f"={lr_sub_cell}"
    ws1[f"E{dir_rows[2]}"] = f"={din_sub_cell}"
    ws1[f"E{dir_rows[3]}"] = f"={mbr_sub_cell}"
    ws1[f"E{dir_rows[4]}"] = f"={b2_sub_cell}"
    ws1[f"E{dir_rows[5]}"] = f"={b3_sub_cell}"
    ws1[f"E{dir_rows[6]}"] = f"={m_sub_cell}+{v_sub_cell}+{u_sub_cell}"
    ws1[f"E{dir_rows[7]}"] = f"={fc_sub_cell}+{el_sub_cell}"
    ws1[f"E{dir_rows[8]}"] = f"={pl_sub_cell}+{pt_sub_cell}+{dc_sub_cell}"
    ws1[f"E{dir_rows[9]}"] = f"={df_sub_cell}+{lu_sub_cell}+{pm_sub_cell}"
    for dr in dir_rows:
        ws1[f"E{dr}"].number_format = FMT_CURRENCY
    ws1[f"E{dir_total_row}"] = f"={total_incl_gst_cell}"

    max_r_sheet1 = r

    # Set Print Area for Sheet 1
    ws1.print_area = f"A1:F{max_r_sheet1}"

    # Apply openpyxl row breaks
    from openpyxl.worksheet.pagebreak import Break
    for pbr in page_break_rows:
        ws1.row_breaks.append(Break(id=pbr - 1, max=6, man=True))

    # ==============================================================================
    # --- SHEET 2: EXECUTIVE SUMMARY DASHBOARD ---
    # ==============================================================================
    ws2 = wb.create_sheet(title="Executive Grand Summary")
    ws2.views.sheetView[0].showGridLines = True

    col_widths_2 = {
        'A': 6,     # Sl No
        'B': 28,    # Section / Room Name
        'C': 16,    # Category / Type
        'D': 12,    # % of Total
        'E': 16,    # Amount Excl. GST (₹)
        'F': 22     # Key Highlights
    }
    for col, width in col_widths_2.items():
        ws2.column_dimensions[col].width = width

    ws2.page_setup.orientation = ws2.ORIENTATION_PORTRAIT
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.page_margins.left = 0.4
    ws2.page_margins.right = 0.4

    r2 = 1
    ws2.row_dimensions[r2].height = 4
    apply_style(ws2, 1, r2, 6, r2, fill=fill_gold_bar)
    r2 += 1

    ws2.row_dimensions[r2].height = 26
    ws2.merge_cells(f"A{r2}:F{r2}")
    ws2[f"A{r2}"] = "DECOR8 INDIA — EXECUTIVE COST BREAKDOWN & PROJECT SUMMARY"
    apply_style(ws2, 1, r2, 6, r2, font=font_main_title, fill=PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid"), alignment=align_center)
    r2 += 1

    ws2.row_dimensions[r2].height = 22
    ws2.merge_cells(f"A{r2}:C{r2}")
    ws2[f"A{r2}"] = "Client: MR. GOVINDA | Project: Kolte Patil D1304"
    apply_style(ws2, 1, r2, 3, r2, font=font_client_hdr, fill=fill_subtotal, alignment=align_left)

    ws2.merge_cells(f"D{r2}:F{r2}")
    ws2[f"D{r2}"] = "Quotation No: D8202607102 | Date: 06-AUG-2026"
    apply_style(ws2, 4, r2, 6, r2, font=font_client_hdr, fill=fill_subtotal, alignment=align_right)
    r2 += 2

    headers_dash = ["#", "Section / Room Description", "Scope Type", "% Total", "Amount (₹)", "Specifications Highlights"]
    ws2.row_dimensions[r2].height = 20
    for c_idx, h_txt in enumerate(headers_dash, start=1):
        c_letter = get_column_letter(c_idx)
        ws2[f"{c_letter}{r2}"] = h_txt
        ws2[f"{c_letter}{r2}"].font = font_tbl_header
        ws2[f"{c_letter}{r2}"].fill = fill_tbl_header
        ws2[f"{c_letter}{r2}"].border = header_border
        ws2[f"{c_letter}{r2}"].alignment = align_center if c_idx in [1, 4] else (align_right if c_idx == 5 else align_left)
    r2 += 1

    dash_start = r2
    dashboard_rows_meta = [
        ("01", "Modular Kitchen (Units)", "Joinery & Woodwork", "Acrylic + Century BWP 710 + Hettich"),
        ("",   "Kitchen Sink & Plumbing Materials", "Sanitary & Fixtures", "Carysil Quartz Sink + Jaquar Mixer"),
        ("02", "Living Room", "Furniture & Panelling", "Vitrified Wall + Rafters + 5-Seater Sofa"),
        ("03", "Dining Area", "Loose Furniture", "6-Seater Marble Dining Table & Chairs"),
        ("04", "Master Bedroom", "Joinery & Furniture", "Lacquered Glass Wardrobe + Queen Bed"),
        ("05", "Bedroom 02", "Joinery & Furniture", "Acrylic Sliding Wardrobe + King Bed"),
        ("06", "Bedroom 03", "Joinery & Furniture", "Acrylic Sliding Wardrobe + King Bed"),
        ("07", "Maid Room", "Joinery & Furniture", "Single Bed with Hydraulic Storage"),
        ("08", "Bathroom Vanity Units", "Sanitary & Joinery", "Marine Carcass + KalingaStone Top + Jaquar"),
        ("09", "Utility Storage", "Joinery & Woodwork", "Base + Wall + Loft Century 710 BWP"),
        ("10", "False Ceiling", "Civil & Ceiling", "Saint-Gobain Gyproc + PVC + Grid"),
        ("11", "Electrical & Lighting", "Electrical", "Havells LED + Chandelier + Profile Lights"),
        ("12", "Plumbing Labour", "Plumbing", "Sanitary, Sink & Pressure Testing"),
        ("13", "Painting", "Finishing", "Asian Paints Royale Luxury (3,850 Sq.ft)"),
        ("14", "Deep Cleaning", "Logistics & Handover", "Professional Deep Cleaning & Sanitization"),
        ("15", "Design Fees", "Architecture & Design", "2D/3D Drawings & Walkthrough Video"),
        ("16", "Loading & Unloading", "Logistics", "Material Transport & Vertical Shifting"),
        ("17", "Project Management", "Supervision & Execution", "Full Tenure Site Oversight & Snagging")
    ]

    for idx, (s_num, s_desc, s_cat, s_high) in enumerate(dashboard_rows_meta):
        ws2.row_dimensions[r2].height = 20
        is_zebra = (idx % 2 == 1)
        fill_curr = fill_zebra if is_zebra else fill_white

        ws2[f"A{r2}"] = s_num
        ws2[f"A{r2}"].font = font_cell_bold
        ws2[f"A{r2}"].fill = fill_curr
        ws2[f"A{r2}"].border = cell_border
        ws2[f"A{r2}"].alignment = align_center

        ws2[f"B{r2}"] = s_desc
        ws2[f"B{r2}"].font = font_cell_bold if s_num else font_cell
        ws2[f"B{r2}"].fill = fill_curr
        ws2[f"B{r2}"].border = cell_border
        ws2[f"B{r2}"].alignment = align_left

        ws2[f"C{r2}"] = s_cat
        ws2[f"C{r2}"].font = font_cell
        ws2[f"C{r2}"].fill = fill_curr
        ws2[f"C{r2}"].border = cell_border
        ws2[f"C{r2}"].alignment = align_left

        sheet1_subtot_cell = section_subtotal_cells[idx][1]
        ws2[f"E{r2}"] = f"='Luxury Proposal (17 Sections)'!{sheet1_subtot_cell}"
        ws2[f"E{r2}"].font = font_cell_amount
        ws2[f"E{r2}"].fill = fill_curr
        ws2[f"E{r2}"].border = cell_border
        ws2[f"E{r2}"].number_format = FMT_CURRENCY
        ws2[f"E{r2}"].alignment = align_right

        total_row_dash = dash_start + len(dashboard_rows_meta)
        ws2[f"D{r2}"] = f"=E{r2}/$E${total_row_dash}"
        ws2[f"D{r2}"].font = font_cell
        ws2[f"D{r2}"].fill = fill_curr
        ws2[f"D{r2}"].border = cell_border
        ws2[f"D{r2}"].number_format = '0.0%'
        ws2[f"D{r2}"].alignment = align_center

        ws2[f"F{r2}"] = s_high
        ws2[f"F{r2}"].font = font_cell
        ws2[f"F{r2}"].fill = fill_curr
        ws2[f"F{r2}"].border = cell_border
        ws2[f"F{r2}"].alignment = align_left

        r2 += 1

    dash_end = r2 - 1

    # Grand Total Row
    ws2.row_dimensions[r2].height = 26
    ws2.merge_cells(f"A{r2}:C{r2}")
    ws2[f"A{r2}"] = "GRAND TOTAL (EXCL. GST)"
    apply_style(ws2, 1, r2, 3, r2, font=font_grand_lbl, fill=fill_grand_total, border=grand_total_border, alignment=align_right)

    ws2[f"D{r2}"] = f"=SUM(D{dash_start}:D{dash_end})"
    ws2[f"D{r2}"].font = font_grand_lbl
    ws2[f"D{r2}"].fill = fill_grand_total
    ws2[f"D{r2}"].border = grand_total_border
    ws2[f"D{r2}"].number_format = '0.0%'
    ws2[f"D{r2}"].alignment = align_center

    ws2[f"E{r2}"] = f"=SUM(E{dash_start}:E{dash_end})"
    ws2[f"E{r2}"].font = font_grand_val
    ws2[f"E{r2}"].fill = fill_grand_total
    ws2[f"E{r2}"].border = grand_total_border
    ws2[f"E{r2}"].number_format = FMT_CURRENCY
    ws2[f"E{r2}"].alignment = align_right

    ws2[f"F{r2}"] = "Base Project Value"
    ws2[f"F{r2}"].font = font_grand_lbl
    ws2[f"F{r2}"].fill = fill_grand_total
    ws2[f"F{r2}"].border = grand_total_border
    ws2[f"F{r2}"].alignment = align_center
    r2 += 1

    # Taxes
    ws2.row_dimensions[r2].height = 20
    ws2.merge_cells(f"A{r2}:D{r2}")
    ws2[f"A{r2}"] = "CGST @ 9%"
    apply_style(ws2, 1, r2, 4, r2, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    ws2[f"E{r2}"] = f"=ROUND(E{r2-1}*0.09, 0)"
    apply_style(ws2, 5, r2, 5, r2, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    ws2[f"E{r2}"].number_format = FMT_CURRENCY
    ws2[f"F{r2}"] = "Central GST (9%)"
    apply_style(ws2, 6, r2, 6, r2, font=font_cell, fill=fill_subtotal, border=cell_border, alignment=align_center)
    r2 += 1

    ws2.row_dimensions[r2].height = 20
    ws2.merge_cells(f"A{r2}:D{r2}")
    ws2[f"A{r2}"] = "SGST @ 9%"
    apply_style(ws2, 1, r2, 4, r2, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    ws2[f"E{r2}"] = f"=ROUND(E{r2-2}*0.09, 0)"
    apply_style(ws2, 5, r2, 5, r2, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    ws2[f"E{r2}"].number_format = FMT_CURRENCY
    ws2[f"F{r2}"] = "State GST (9%)"
    apply_style(ws2, 6, r2, 6, r2, font=font_cell, fill=fill_subtotal, border=cell_border, alignment=align_center)
    r2 += 1

    ws2.row_dimensions[r2].height = 24
    ws2.merge_cells(f"A{r2}:D{r2}")
    ws2[f"A{r2}"] = "FINAL ESTIMATED PROJECT VALUE (INCL. 18% GST)"
    apply_style(ws2, 1, r2, 4, r2, font=Font(name="Segoe UI", size=10, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)
    ws2[f"E{r2}"] = f"=E{r2-3}+E{r2-2}+E{r2-1}"
    apply_style(ws2, 5, r2, 5, r2, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)
    ws2[f"E{r2}"].number_format = FMT_CURRENCY
    ws2[f"F{r2}"] = "Complete Turnkey Cost"
    apply_style(ws2, 6, r2, 6, r2, font=font_cell_bold, fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_center)
    r2 += 2

    ws2.row_dimensions[r2].height = 20
    ws2.merge_cells(f"A{r2}:F{r2}")
    ws2[f"A{r2}"] = "Payment Milestones: 15% Booking Confirmation  |  50% Factory Production  |  35% Pre-Dispatch to Site"
    apply_style(ws2, 1, r2, 6, r2, font=font_cell_bold, fill=fill_meta_lbl, border=cell_border, alignment=align_center)

    ws2.print_area = f"A1:F{r2}"

    # ==============================================================================
    # --- SHEET 3: GOVINDA D1304 (ORIGINAL BOQ FORMAT ELEVATED) ---
    # ==============================================================================
    ws3 = wb.create_sheet(title="Govinda D1304 (BOQ Format)")
    ws3.views.sheetView[0].showGridLines = True

    for col, width in col_widths_1.items():
        ws3.column_dimensions[col].width = width

    ws3.page_setup.orientation = ws3.ORIENTATION_PORTRAIT
    ws3.page_setup.paperSize = ws3.PAPERSIZE_A4
    ws3.sheet_properties.pageSetUpPr.fitToPage = True
    ws3.page_setup.fitToWidth = 1
    ws3.page_setup.fitToHeight = 0
    ws3.page_margins.left = 0.4
    ws3.page_margins.right = 0.4

    r3 = 1
    ws3.row_dimensions[r3].height = 4
    apply_style(ws3, 1, r3, 6, r3, fill=fill_gold_bar)
    r3 += 1

    ws3.row_dimensions[r3].height = 26
    ws3.merge_cells(f"A{r3}:F{r3}")
    ws3[f"A{r3}"] = "DECOR8 INDIA — BILL OF QUANTITIES (BOQ)"
    apply_style(ws3, 1, r3, 6, r3, font=font_main_title, fill=PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid"), alignment=align_center)
    r3 += 1

    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:C{r3}")
    ws3[f"A{r3}"] = "QUOTATION FOR — MR. GOVINDA - Kolte Patil D1304"
    apply_style(ws3, 1, r3, 3, r3, font=font_client_hdr, fill=fill_subtotal, alignment=align_left)

    ws3.merge_cells(f"D{r3}:F{r3}")
    ws3[f"D{r3}"] = "Representative: Mr. Satish Bhat"
    apply_style(ws3, 4, r3, 6, r3, font=font_client_hdr, fill=fill_subtotal, alignment=align_right)
    r3 += 2

    # Section A
    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:F{r3}")
    ws3[f"A{r3}"] = "Section A: Wood work & Modular Finish CENTURY PLY"
    apply_style(ws3, 1, r3, 6, r3, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE), fill=fill_tbl_header, border=header_border, alignment=align_left)
    r3 += 1

    headers_boq = ["Item Name", "Sub Item Name / Description", "Total SQFT", "Unit", "Rate", "Amount"]
    ws3.row_dimensions[r3].height = 20
    for c_idx, h_txt in enumerate(headers_boq, start=1):
        c_letter = get_column_letter(c_idx)
        ws3[f"{c_letter}{r3}"] = h_txt
        ws3[f"{c_letter}{r3}"].font = font_tbl_header
        ws3[f"{c_letter}{r3}"].fill = fill_tbl_header
        ws3[f"{c_letter}{r3}"].border = header_border
        ws3[f"{c_letter}{r3}"].alignment = align_center if c_idx in [3, 4] else (align_right if c_idx in [5, 6] else align_left)
    r3 += 1

    sec_a_start = r3
    boq_a_items = [
        ("Utility storage", "Profile shutters — For TV unit and kitchen", 26, "Sq.ft", 1100),
        ("Loft", "Loft — Frame with shutters", 38, "Sq.ft", 950),
        ("Tall Unit", "Tall Storage Unit", 18, "Sq.ft", 1550),
        ("Wicker Basket", "Wicker Basket — 2 Baskets", 1, "Set", 12000),
        ("KBR Wardrobe", "Swing Wardrobe — Century MR Ply carcase & Laminate Finish", 60, "Sq.ft", 1550),
        ("KBR Loft", "Loft — Frame with shutters", 36, "Sq.ft", 950),
    ]
    for idx, (item, sub, sqft, unit, rate) in enumerate(boq_a_items):
        ws3.row_dimensions[r3].height = 20
        fill_curr = fill_zebra if (idx % 2 == 1) else fill_white
        ws3[f"A{r3}"] = item
        ws3[f"A{r3}"].font = font_cell_bold
        ws3[f"A{r3}"].fill = fill_curr
        ws3[f"A{r3}"].border = cell_border
        ws3[f"A{r3}"].alignment = align_left

        ws3[f"B{r3}"] = sub
        ws3[f"B{r3}"].font = font_cell
        ws3[f"B{r3}"].fill = fill_curr
        ws3[f"B{r3}"].border = cell_border
        ws3[f"B{r3}"].alignment = align_left

        ws3[f"C{r3}"] = sqft
        ws3[f"C{r3}"].font = font_cell
        ws3[f"C{r3}"].fill = fill_curr
        ws3[f"C{r3}"].border = cell_border
        ws3[f"C{r3}"].alignment = align_center

        ws3[f"D{r3}"] = unit
        ws3[f"D{r3}"].font = font_cell
        ws3[f"D{r3}"].fill = fill_curr
        ws3[f"D{r3}"].border = cell_border
        ws3[f"D{r3}"].alignment = align_center

        ws3[f"E{r3}"] = rate
        ws3[f"E{r3}"].font = font_cell
        ws3[f"E{r3}"].fill = fill_curr
        ws3[f"E{r3}"].border = cell_border
        ws3[f"E{r3}"].number_format = FMT_CURRENCY
        ws3[f"E{r3}"].alignment = align_right

        ws3[f"F{r3}"] = f"=C{r3}*E{r3}"
        ws3[f"F{r3}"].font = font_cell_amount
        ws3[f"F{r3}"].fill = fill_curr
        ws3[f"F{r3}"].border = cell_border
        ws3[f"F{r3}"].number_format = FMT_CURRENCY
        ws3[f"F{r3}"].alignment = align_right

        r3 += 1

    sec_a_subtot_row = r3
    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "Sub Total excluding GST (A)"
    apply_style(ws3, 1, r3, 5, r3, font=font_cell_subtot, fill=fill_subtotal, border=subtotal_border, alignment=align_right)

    ws3[f"F{r3}"] = f"=SUM(F{sec_a_start}:F{sec_a_subtot_row-1})"
    ws3[f"F{r3}"].font = font_cell_subtot
    ws3[f"F{r3}"].fill = fill_subtotal
    ws3[f"F{r3}"].border = subtotal_border
    ws3[f"F{r3}"].number_format = FMT_CURRENCY
    ws3[f"F{r3}"].alignment = align_right
    r3 += 2

    # Section B: Appliances & Accessories
    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:F{r3}")
    ws3[f"A{r3}"] = "Section B: Appliances & Accessories"
    apply_style(ws3, 1, r3, 6, r3, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE), fill=fill_tbl_header, border=header_border, alignment=align_left)
    r3 += 1

    ws3.row_dimensions[r3].height = 20
    for c_idx, h_txt in enumerate(headers_boq, start=1):
        c_letter = get_column_letter(c_idx)
        ws3[f"{c_letter}{r3}"] = h_txt
        ws3[f"{c_letter}{r3}"].font = font_tbl_header
        ws3[f"{c_letter}{r3}"].fill = fill_tbl_header
        ws3[f"{c_letter}{r3}"].border = header_border
        ws3[f"{c_letter}{r3}"].alignment = align_center if c_idx in [3, 4] else (align_right if c_idx in [5, 6] else align_left)
    r3 += 1

    sec_b_start = r3
    boq_b_items = [
        ("Kitchen Hardware", "Hob and Chimney (Elica / Kutchina)", 1, "Set", 0),
        ("Kitchen Hardware", "Sink", 1, "No.", 0),
        ("Kitchen Hardware", "Inbuilt Oven and Microwave (Model selection price may vary)", 1, "Set", 0),
        ("Kitchen Hardware", "Copper pipe for Gas Cylinder connection", 1, "L.S.", 0),
        ("Kitchen Hardware", "1x cutlery with tandem drawer", 1, "No.", 0),
        ("Kitchen Hardware", "1x plain ss basket drawer", 1, "No.", 0),
        ("Kitchen Hardware", "Rolling shutter", 1, "No.", 0),
        ("Kitchen Hardware", "1x thali basket, Magic corner, Wicker basket", 1, "Set", 0),
        ("Wardrobes Hardware", "Handles, hinges and required hardwares added", 1, "Set", 23500),
    ]
    for idx, (item, sub, sqft, unit, rate) in enumerate(boq_b_items):
        ws3.row_dimensions[r3].height = 20
        fill_curr = fill_zebra if (idx % 2 == 1) else fill_white
        ws3[f"A{r3}"] = item
        ws3[f"A{r3}"].font = font_cell_bold
        ws3[f"A{r3}"].fill = fill_curr
        ws3[f"A{r3}"].border = cell_border
        ws3[f"A{r3}"].alignment = align_left

        ws3[f"B{r3}"] = sub
        ws3[f"B{r3}"].font = font_cell
        ws3[f"B{r3}"].fill = fill_curr
        ws3[f"B{r3}"].border = cell_border
        ws3[f"B{r3}"].alignment = align_left

        ws3[f"C{r3}"] = sqft
        ws3[f"C{r3}"].font = font_cell
        ws3[f"C{r3}"].fill = fill_curr
        ws3[f"C{r3}"].border = cell_border
        ws3[f"C{r3}"].alignment = align_center

        ws3[f"D{r3}"] = unit
        ws3[f"D{r3}"].font = font_cell
        ws3[f"D{r3}"].fill = fill_curr
        ws3[f"D{r3}"].border = cell_border
        ws3[f"D{r3}"].alignment = align_center

        ws3[f"E{r3}"] = rate
        ws3[f"E{r3}"].font = font_cell
        ws3[f"E{r3}"].fill = fill_curr
        ws3[f"E{r3}"].border = cell_border
        ws3[f"E{r3}"].number_format = FMT_CURRENCY
        ws3[f"E{r3}"].alignment = align_right

        ws3[f"F{r3}"] = f"=C{r3}*E{r3}"
        ws3[f"F{r3}"].font = font_cell_amount
        ws3[f"F{r3}"].fill = fill_curr
        ws3[f"F{r3}"].border = cell_border
        ws3[f"F{r3}"].number_format = FMT_CURRENCY
        ws3[f"F{r3}"].alignment = align_right

        r3 += 1

    sec_b_subtot_row = r3
    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "Sub Total excluding GST (B)"
    apply_style(ws3, 1, r3, 5, r3, font=font_cell_subtot, fill=fill_subtotal, border=subtotal_border, alignment=align_right)

    ws3[f"F{r3}"] = f"=SUM(F{sec_b_start}:F{sec_b_subtot_row-1})"
    ws3[f"F{r3}"].font = font_cell_subtot
    ws3[f"F{r3}"].fill = fill_subtotal
    ws3[f"F{r3}"].border = subtotal_border
    ws3[f"F{r3}"].number_format = FMT_CURRENCY
    ws3[f"F{r3}"].alignment = align_right
    r3 += 2

    # Section C: Miscellaneous & Loose Furniture
    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:F{r3}")
    ws3[f"A{r3}"] = "Section C: Miscellaneous & Loose Furniture"
    apply_style(ws3, 1, r3, 6, r3, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE), fill=fill_tbl_header, border=header_border, alignment=align_left)
    r3 += 1

    ws3.row_dimensions[r3].height = 20
    for c_idx, h_txt in enumerate(headers_boq, start=1):
        c_letter = get_column_letter(c_idx)
        ws3[f"{c_letter}{r3}"] = h_txt
        ws3[f"{c_letter}{r3}"].font = font_tbl_header
        ws3[f"{c_letter}{r3}"].fill = fill_tbl_header
        ws3[f"{c_letter}{r3}"].border = header_border
        ws3[f"{c_letter}{r3}"].alignment = align_center if c_idx in [3, 4] else (align_right if c_idx in [5, 6] else align_left)
    r3 += 1

    sec_c_start = r3
    boq_c_items = [
        ("King size bed", "King size bed Hydraulic with Headboard and side table", 1, "Unit", 0),
        ("Queen size bed", "Queen size bed Hydraulic with Headboard and side table with cushion", 1, "Unit", 0),
        ("Wallpapers", "Premium textured designer wallpapers as per selection", 1, "L.S.", 0),
        ("Painting", "Royale Aspira — 1 coat primer + 2 coats Asian Paints Aspira", 1, "L.S.", 0),
        ("Quartz/Granite stone", "Countertop stone including cutting, edge moulding & installation", 1, "L.S.", 0),
        ("Electrical work", "Wiring, pop wiring, switches, modular plates etc.", 1, "L.S.", 0),
    ]
    for idx, (item, sub, sqft, unit, rate) in enumerate(boq_c_items):
        ws3.row_dimensions[r3].height = 20
        fill_curr = fill_zebra if (idx % 2 == 1) else fill_white
        ws3[f"A{r3}"] = item
        ws3[f"A{r3}"].font = font_cell_bold
        ws3[f"A{r3}"].fill = fill_curr
        ws3[f"A{r3}"].border = cell_border
        ws3[f"A{r3}"].alignment = align_left

        ws3[f"B{r3}"] = sub
        ws3[f"B{r3}"].font = font_cell
        ws3[f"B{r3}"].fill = fill_curr
        ws3[f"B{r3}"].border = cell_border
        ws3[f"B{r3}"].alignment = align_left

        ws3[f"C{r3}"] = sqft
        ws3[f"C{r3}"].font = font_cell
        ws3[f"C{r3}"].fill = fill_curr
        ws3[f"C{r3}"].border = cell_border
        ws3[f"C{r3}"].alignment = align_center

        ws3[f"D{r3}"] = unit
        ws3[f"D{r3}"].font = font_cell
        ws3[f"D{r3}"].fill = fill_curr
        ws3[f"D{r3}"].border = cell_border
        ws3[f"D{r3}"].alignment = align_center

        ws3[f"E{r3}"] = rate
        ws3[f"E{r3}"].font = font_cell
        ws3[f"E{r3}"].fill = fill_curr
        ws3[f"E{r3}"].border = cell_border
        ws3[f"E{r3}"].number_format = FMT_CURRENCY
        ws3[f"E{r3}"].alignment = align_right

        ws3[f"F{r3}"] = f"=C{r3}*E{r3}"
        ws3[f"F{r3}"].font = font_cell_amount
        ws3[f"F{r3}"].fill = fill_curr
        ws3[f"F{r3}"].border = cell_border
        ws3[f"F{r3}"].number_format = FMT_CURRENCY
        ws3[f"F{r3}"].alignment = align_right

        r3 += 1

    sec_c_subtot_row = r3
    ws3.row_dimensions[r3].height = 22
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "Sub Total excluding GST (C)"
    apply_style(ws3, 1, r3, 5, r3, font=font_cell_subtot, fill=fill_subtotal, border=subtotal_border, alignment=align_right)

    ws3[f"F{r3}"] = f"=SUM(F{sec_c_start}:F{sec_c_subtot_row-1})"
    ws3[f"F{r3}"].font = font_cell_subtot
    ws3[f"F{r3}"].fill = fill_subtotal
    ws3[f"F{r3}"].border = subtotal_border
    ws3[f"F{r3}"].number_format = FMT_CURRENCY
    ws3[f"F{r3}"].alignment = align_right
    r3 += 2

    # Grand Total A+B+C
    ws3.row_dimensions[r3].height = 26
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "Total Project Cost (A+B+C) excluding GST"
    apply_style(ws3, 1, r3, 5, r3, font=font_grand_lbl, fill=fill_grand_total, border=grand_total_border, alignment=align_right)

    boq_tot_excl = f"F{r3}"
    ws3[boq_tot_excl] = f"=F{sec_a_subtot_row}+F{sec_b_subtot_row}+F{sec_c_subtot_row}"
    ws3[boq_tot_excl].font = font_grand_val
    ws3[boq_tot_excl].fill = fill_grand_total
    ws3[boq_tot_excl].border = grand_total_border
    ws3[boq_tot_excl].number_format = FMT_CURRENCY
    ws3[boq_tot_excl].alignment = align_right
    r3 += 1

    # Taxes
    ws3.row_dimensions[r3].height = 20
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "CGST @ 9%"
    apply_style(ws3, 1, r3, 5, r3, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    boq_cgst = f"F{r3}"
    ws3[boq_cgst] = f"=ROUND({boq_tot_excl}*0.09, 0)"
    apply_style(ws3, 6, r3, 6, r3, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    ws3[boq_cgst].number_format = FMT_CURRENCY
    r3 += 1

    ws3.row_dimensions[r3].height = 20
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "SGST @ 9%"
    apply_style(ws3, 1, r3, 5, r3, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    boq_sgst = f"F{r3}"
    ws3[boq_sgst] = f"=ROUND({boq_tot_excl}*0.09, 0)"
    apply_style(ws3, 6, r3, 6, r3, font=font_cell_bold, fill=fill_subtotal, border=cell_border, alignment=align_right)
    ws3[boq_sgst].number_format = FMT_CURRENCY
    r3 += 1

    ws3.row_dimensions[r3].height = 24
    ws3.merge_cells(f"A{r3}:E{r3}")
    ws3[f"A{r3}"] = "Total Project Cost (A+B+C) including GST"
    apply_style(ws3, 1, r3, 5, r3, font=Font(name="Segoe UI", size=10, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)
    boq_tot_incl = f"F{r3}"
    ws3[boq_tot_incl] = f"={boq_tot_excl}+{boq_cgst}+{boq_sgst}"
    apply_style(ws3, 6, r3, 6, r3, font=Font(name="Segoe UI", size=11, bold=True, color=COLOR_DARK), fill=PatternFill(start_color="EBD6A7", end_color="EBD6A7", fill_type="solid"), border=grand_total_border, alignment=align_right)
    ws3[boq_tot_incl].number_format = FMT_CURRENCY
    r3 += 2

    ws3.print_area = f"A1:F{r3}"

    # Save initial openpyxl files
    for p in output_excel_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        wb.save(p)
        print(f"Saved initial workbook: {p}")

    # ==============================================================================
    # --- EXCEL COM POST-PROCESSING: PERFECT PAGE SETUP & BORDERS ---
    # ==============================================================================
    excel = None
    try:
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        for path in output_excel_paths:
            abs_p = os.path.abspath(path)
            xl_wb = excel.Workbooks.Open(abs_p)
            excel.Calculate()

            for sheet_name in ["Luxury Proposal (17 Sections)", "Executive Grand Summary", "Govinda D1304 (BOQ Format)"]:
                try:
                    xl_ws = xl_wb.Sheets(sheet_name)
                    # Force exact A4 Portrait, Fit to 1 Page Wide!
                    xl_ws.PageSetup.Orientation = 1  # 1 = xlPortrait
                    xl_ws.PageSetup.PaperSize = 9    # 9 = xlPaperA4
                    xl_ws.PageSetup.Zoom = False
                    xl_ws.PageSetup.FitToPagesWide = 1
                    xl_ws.PageSetup.FitToPagesTall = False
                    xl_ws.PageSetup.LeftMargin = excel.InchesToPoints(0.25)
                    xl_ws.PageSetup.RightMargin = excel.InchesToPoints(0.25)
                    xl_ws.PageSetup.TopMargin = excel.InchesToPoints(0.25)
                    xl_ws.PageSetup.BottomMargin = excel.InchesToPoints(0.25)
                    xl_ws.PageSetup.CenterHorizontally = True

                    if sheet_name == "Luxury Proposal (17 Sections)":
                        for pbr in page_break_rows:
                            try:
                                xl_ws.HPageBreaks.Add(Before=xl_ws.Rows(pbr))
                            except Exception:
                                pass
                except Exception as ps_err:
                    print(f"PageSetup note for {sheet_name}: {ps_err}")

            xl_wb.Save()
            xl_wb.Close()
            print(f"Excel COM PageSetup verified: {abs_p}")

        print("Excel COM post-processing finished successfully.")
    except Exception as com_err:
        print(f"COM post-processing error: {com_err}")
    finally:
        if excel:
            try:
                excel.Quit()
            except Exception:
                pass

if __name__ == "__main__":
    build_perfect_white_gold_quotation()
